import tempfile
import os
from typing import Optional, List, Any

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from vectorstore_manager import get_vectorstore_for_chat


def parse_file_to_text(file: Any, mime_type: Optional[str] = None) -> str:
	"""
	Parses an uploaded file-like object (e.g. Flask FileStorage) to plain text.
	Supports txt, pdf, docx, xlsx. Returns extracted string.
	"""
	mime_type = mime_type or getattr(file, 'mimetype', '') or ''
	filename = getattr(file, 'filename', '') or 'tempfile'
	with tempfile.NamedTemporaryFile(delete=True, suffix=os.path.splitext(filename)[1]) as tmp:
		# Prefer .save if available; fall back to writing bytes from .read
		if hasattr(file, 'save'):
			file.save(tmp.name)
		else:
			with open(tmp.name, 'wb') as out:
				out.write(file.read())
		if mime_type.startswith('text/plain') or filename.endswith('.txt'):
			with open(tmp.name, 'r', encoding='utf-8', errors='ignore') as f:
				return f.read()
		elif mime_type == 'application/pdf' or filename.endswith('.pdf'):
			import pdfplumber
			with pdfplumber.open(tmp.name) as pdf:
				pages = [page.extract_text() or '' for page in pdf.pages]
			return '\n'.join(pages)
		elif (mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' or filename.endswith('.docx')):
			from docx import Document as DocxDocument
			doc = DocxDocument(tmp.name)
			paragraphs = [p.text for p in doc.paragraphs]
			return '\n'.join(paragraphs)
		elif (mime_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' or filename.endswith('.xlsx')):
			import pandas as pd
			df = pd.read_excel(tmp.name)
			return df.to_csv(index=False)
		else:
			return ""


def ingest_text_to_vectorstore(chat_id: str, text: str, source_name: Optional[str] = None) -> None:
	if not str(text).strip():
		return
	text_splitter = RecursiveCharacterTextSplitter(chunk_size=250, chunk_overlap=50)
	chunks = text_splitter.split_text(text)
	docs = [Document(page_content=c, metadata={'source': source_name or 'uploaded'}) for c in chunks]
	if docs:
		get_vectorstore_for_chat(chat_id).add_documents(docs)


def ingest_files(chat_id: str, files: List[Any]) -> None:
	if not files:
		return
	for f in files:
		mime = getattr(f, 'mimetype', None)
		name = getattr(f, 'filename', None)
		text = parse_file_to_text(f, mime)
		ingest_text_to_vectorstore(chat_id, text, name)


# Debug/testing helpers
def list_chat_collection_name(chat_id: str) -> str:
    from vectorstore_manager import _collection_name_for_chat
    return _collection_name_for_chat(chat_id)

def list_ingested_files_for_chat(chat_id: str) -> List[str]:
    """
    Returns a list of distinct source filenames stored as metadata in the vector store for a chat.
    Note: This relies on PGVector's underlying table storing metadata JSON. If not available,
    this function will return an empty list.
    """
    try:
        from database import db_connection
        collection = list_chat_collection_name(chat_id)
        with db_connection.cursor() as cur:
            # Assuming a table naming convention used by PGVector for collections
            # Fallback: try a generic metadata query from langchain_pg_embedding table
            cur.execute("""
                SELECT DISTINCT metadata->>'source' AS source
                FROM langchain_pg_embedding
                WHERE collection_id = (
                    SELECT uuid FROM langchain_pg_collection WHERE name = %s
                ) AND metadata ? 'source'
            """, (collection,))
            rows = cur.fetchall()
            return [r[0] for r in rows if r and r[0]]
    except Exception:
        return []

