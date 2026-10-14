"""File parsing, chunking, embeddings and the pgvector store."""
import os

from fastembed import TextEmbedding
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_postgres import Column, PGEngine, PGVectorStore
from langchain_text_splitters import RecursiveCharacterTextSplitter

from backend import config

TABLE = "chunks"
SNIPPETS = "highlights"  # source name for a folder's highlighted snippets

_splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=100)
_store = None


class FastEmbedEmbeddings(Embeddings):
    """Local ONNX embeddings (no torch)."""

    def __init__(self, model_name):
        self.model = TextEmbedding(model_name)

    def embed_documents(self, texts):
        return [v.tolist() for v in self.model.embed(texts)]

    def embed_query(self, text):
        return self.embed_documents([text])[0]


def store():
    global _store
    if _store is None:
        embeddings = FastEmbedEmbeddings(config.EMBED_MODEL)
        url = config.DATABASE_URL.replace("postgresql://", "postgresql+psycopg://", 1)
        engine = PGEngine.from_connection_string(url)
        try:
            engine.init_vectorstore_table(
                TABLE,
                len(embeddings.embed_query("dim")),
                metadata_columns=[Column("folder_id", "TEXT"), Column("source", "TEXT")],
            )
        except Exception as e:
            if "already exists" not in str(e):
                raise
        _store = PGVectorStore.create_sync(
            engine, embeddings, TABLE, metadata_columns=["folder_id", "source"]
        )
    return _store


def parse_file(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".txt":
        with open(path, encoding="utf-8", errors="ignore") as f:
            return f.read()
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if ext == ".docx":
        from docx import Document as Docx
        return "\n".join(p.text for p in Docx(path).paragraphs)
    if ext == ".xlsx":
        from openpyxl import load_workbook
        lines = []
        for sheet in load_workbook(path, read_only=True, data_only=True):
            for row in sheet.iter_rows(values_only=True):
                lines.append(",".join("" if v is None else str(v) for v in row))
        return "\n".join(lines)
    raise ValueError(f"Unsupported file type: {ext}")


def index(folder_id, source, texts):
    """Replace a source's chunks in a folder. `texts` are split into chunks before embedding."""
    unindex(folder_id, source)
    docs = [
        Document(page_content=chunk, metadata={"folder_id": folder_id, "source": source})
        for text in texts
        for chunk in _splitter.split_text(text)
    ]
    if docs:
        store().add_documents(docs)


def unindex(folder_id, source=None):
    filter = {"folder_id": folder_id}
    if source:
        filter["source"] = source
    store().delete(filter=filter)


def search(folder_id, query, k=4):
    return store().similarity_search(query, k=k, filter={"folder_id": folder_id})
